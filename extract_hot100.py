"""Extract supplied handbook content without executing any embedded code."""
import ast
import json
import re
from pathlib import Path
from urllib.parse import urlparse
from docx import Document
from docx.oxml.ns import qn

import sys
if len(sys.argv)!=2: raise SystemExit('Usage: python extract_hot100.py SOURCE.docx')
SOURCE = Path(sys.argv[1])

def extract():
    doc = Document(SOURCE)
    topics, questions = [], []
    topic = question = None
    guide = False
    for p in doc.paragraphs:
        text, style = p.text, p.style.name
        if style == 'Heading 1':
            match = re.fullmatch(r'专题\s+(\d+)\s+(.+)', text)
            question = None
            guide = bool(match)
            topic = {'name': match[2], 'guide': []} if match else None
            if topic:
                topics.append(topic)
            continue
        match = re.fullmatch(r'第\s+(\d+)\s+题\s+(\d+)\.\s+(.+)', text) if style == 'Heading 2' else None
        if match:
            assert topic is not None
            question = {'id': 'lc-' + match[2], 'order': int(match[1]), 'number': int(match[2]),
                        'title': match[3], 'topic': len(topics)-1, 'blocks': []}
            questions.append(question)
            guide = False
            continue
        if not text.strip():
            continue
        if question is not None:
            if style == 'Problem Meta':
                meta = text.split('|')
                question['difficulty'], question['pattern'] = meta[1].strip(), meta[2].strip()
                links = [doc.part.rels[h.get(qn('r:id'))].target_ref for h in p._p.findall(qn('w:hyperlink'))]
                assert len(links) == 1
                link = urlparse(links[0])
                assert link.scheme == 'https' and link.hostname == 'leetcode.cn' and link.path.startswith('/problems/')
                question['url'] = links[0]
            elif text.startswith('题意概括：'):
                question['stem'] = text.removeprefix('题意概括：')
            else:
                question['blocks'].append({'style': style, 'text': text})
        elif guide and topic is not None:
            if text == '本专题题目地图':
                guide = False
            else:
                topic['guide'].append({'style': style, 'text': text})
    assert len(topics) == 17
    assert [q['order'] for q in questions] == list(range(1, 101))
    assert len({q['id'] for q in questions}) == 100
    # Each topic's original table independently verifies title, number and difficulty.
    table_rows = [row for table in doc.tables[2:] for row in table.rows[1:]]
    expected = [(int(r.cells[0].text), r.cells[1].text, r.cells[2].text) for r in table_rows]
    assert expected == [(q['number'], q['title'], q['difficulty']) for q in questions]
    for q in questions:
        assert all(q.get(key) for key in ['stem','difficulty','pattern','url','blocks'])
        labels = [b['text'] for b in q['blocks'] if b['style'] == 'Section Label']
        assert len(labels) == 7, (q['id'], labels)
        code_blocks = [b['text'] for b in q['blocks'] if b['style'] == 'Code Block']
        assert len(code_blocks) == 2, q['id']
        for code in code_blocks:
            ast.parse(code)  # Syntax only; does not execute document content.
    result = {'topics': topics, 'questions': questions, 'source': SOURCE.name}
    Path(__file__).with_name('hot100.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    print('Validated Hot100: 100 unique questions, 17 topics, 200 Python code blocks; all titles/numbers/difficulties match the document tables.')
    return result

if __name__ == '__main__':
    extract()
